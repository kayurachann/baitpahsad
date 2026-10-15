"""Định dạng chung theo bản cũ "Net-Viet-Travel-De-an-CHOT-CUOI" (đo từ PDF bản cũ).

A4; lề trái 3 cm, phải 2 cm, trên và dưới 2 cm; Times New Roman 13 pt; giãn dòng 1,5;
căn đều; thụt đầu dòng 1 cm; tiêu đề in đậm cùng cỡ chữ; chú thích bảng 12 pt in đậm ở trên bảng;
chữ trong bảng 12 pt, hàng tiêu đề nền xám; số trang ở giữa chân trang, trang bìa không đánh số.
"""
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
TEXT_WIDTH_CM = 21.0 - 3.0 - 2.0


def _set_font(style_or_run, size=None, bold=None, italic=None):
    font = style_or_run.font
    font.name = FONT
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    font.color.rgb = RGBColor(0, 0, 0)


def new_document(title="", subject="", author=""):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

    st = doc.styles
    normal = st["Normal"]
    _set_font(normal, 13)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.0)
    pf.widow_control = True

    for name, size, before, after in (("Heading 1", 13, 12, 6), ("Heading 2", 13, 8, 3)):
        h = st[name]
        _set_font(h, size, bold=True, italic=False)
        hf = h.paragraph_format
        hf.space_before, hf.space_after = Pt(before), Pt(after)
        hf.first_line_indent = Cm(0)
        hf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hf.keep_with_next = True
        hf.line_spacing = 1.5

    cap = st["Caption"]
    _set_font(cap, 12, bold=True, italic=False)
    cf = cap.paragraph_format
    cf.space_before, cf.space_after = Pt(6), Pt(3)
    cf.first_line_indent = Cm(0)
    cf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cf.keep_with_next = True
    cf.line_spacing = 1.0

    core = doc.core_properties
    core.title, core.subject, core.author = title, subject, author
    return doc


def add_runs(par, text, size=None, bold=False, italic=False):
    """Thêm chữ có **in đậm** và *in nghiêng* vào đoạn."""
    for chunk in re.split(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not chunk:
            continue
        b, i, t = bold, italic, chunk
        if chunk.startswith("**") and chunk.endswith("**"):
            b, t = True, chunk[2:-2]
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            i, t = True, chunk[1:-1]
        run = par.add_run(t)
        _set_font(run, size, bold=b, italic=i)
    return par


def para(doc_or_cell, text="", lead=None, indent=True, align=None, size=None, bold=False, italic=False,
         space_after=None, space_before=None, line_spacing=None, keep_next=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    if not indent:
        pf.first_line_indent = Cm(0)
    if align is not None:
        pf.alignment = align
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    if keep_next:
        pf.keep_with_next = True
    if lead:
        add_runs(p, lead.rstrip() + " ", size=size, bold=True, italic=italic)
    add_runs(p, text, size=size, bold=bold, italic=italic)
    return p


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_runs(p, text, size=13, bold=True)
    return p


def caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    add_runs(p, text, size=12, bold=True)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(table, sz=4):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblpr.append(borders)


def _cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def _cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def table(doc, header, rows, widths=None, numeric_cols=(), bold_first_col=None, size=12):
    ncol = len(header)
    widths = widths or [100 / ncol] * ncol
    total = sum(widths)
    widths_cm = [TEXT_WIDTH_CM * w / total for w in widths]
    if bold_first_col is None:
        bold_first_col = not numeric_cols and ncol <= 3

    t = doc.add_table(rows=1 + len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    _cell_margins(t)

    def fill(cell, text, bold=False, right=False, keep=False):
        p = cell.paragraphs[0]
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT
        if keep:
            pf.keep_with_next = True
        lines = str(text).split("\n")
        for k, line in enumerate(lines):
            if k:
                p = cell.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.alignment = pf.alignment
            add_runs(p, line, size=size, bold=bold)

    hdr = t.rows[0]
    _repeat_header(hdr)
    for c, h in enumerate(header):
        cell = hdr.cells[c]
        _shade(cell, "D9D9D9")
        fill(cell, h.replace("**", ""), bold=True, right=c in numeric_cols, keep=True)
    for r, row in enumerate(rows, start=1):
        trow = t.rows[r]
        _cant_split(trow)
        is_total = bool(row) and str(row[0]).strip().startswith("**") and str(row[0]).strip().endswith("**")
        for c in range(ncol):
            text = row[c] if c < len(row) else ""
            clean = str(text)
            bold = is_total or (c == 0 and bold_first_col)
            if bold:
                clean = clean.replace("**", "")
            fill(trow.cells[c], clean, bold=bold, right=c in numeric_cols, keep=r < len(rows) and len(rows) <= 4)
    for row in t.rows:
        for c, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[c])
    # LibreOffice đọc bề rộng cột từ tblGrid và tblW, không đọc từ từng ô
    twips = [int(round(w / 2.54 * 1440)) for w in widths_cm]
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), twips):
        gc.set(qn("w:w"), str(w))
    tblpr = t._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(twips)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    # khoảng cách sau bảng
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 0.8
    return t


def steps(doc, items, size=13):
    for it in items:
        m = re.match(r"^(Bước\s+\d+\.)\s*(.*)$", it)
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(-2.0)
        pf.left_indent = Cm(2.6)
        pf.space_after = Pt(2)
        pf.tab_stops.add_tab_stop(Cm(2.6))
        if m:
            add_runs(p, m.group(1) + "\t" + m.group(2), size=size)
        else:
            add_runs(p, it, size=size)


def note(doc, text):
    p = para(doc, text, indent=False, size=11, italic=True, space_after=6, line_spacing=1.15)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def ref(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(-0.8)
    pf.left_indent = Cm(0.8)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    add_runs(p, text, size=12)
    return p


def page_number_footer(doc, skip_first=True):
    sec = doc.sections[0]
    if skip_first:
        sec.different_first_page_header_footer = True
    p = sec.footer.paragraphs[0]
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    _set_font(run, 12)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)


def toc_line(doc, text, page, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_after = Pt(4)
    pf.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    add_runs(p, f"{text}\t{page}", size=13, bold=bold)
    return p

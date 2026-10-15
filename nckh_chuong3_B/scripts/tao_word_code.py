"""Tạo ch3_B_code_colab.docx: code Google Colab của người B, lấy nguyên văn từ B_3_5_1_3_5_2.ipynb.

Mỗi ô code kèm mục dùng nó và kết quả mong đợi (số của lần chạy thật trên tệp chính; số của tệp Original
theo kết quả chạy thử trong file kế hoạch).
Cách chạy: python3 tao_word_code.py
"""
import json
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import tao_word as tw  # dùng lại định dạng trang và các hàm doan, bang, tieu_de

GOC = Path(__file__).resolve().parent.parent
nb = json.loads((GOC / "B_3_5_1_3_5_2.ipynb").read_text(encoding="utf-8"))
O_CODE = ["".join(o["source"]) for o in nb["cells"] if o["cell_type"] == "code"]
assert len(O_CODE) == 7, len(O_CODE)

# (tiêu đề, mục, kết quả mong đợi) theo đúng thứ tự 7 ô code trong notebook
MO_TA = [
    ("Ô 1. Khởi tạo (chạy trước mọi ô khác)", "Dùng chung cho 3.5.1 và 3.5.2",
     "Dòng cuối in ra (996, 28) (791, 28). Nếu in ra (996, 28) (996, 28) là đang đọc nhầm tệp Original."),
    ("Ô 2. Đối chiếu mã hóa với Bảng 3.1", "Mục 3.3",
     "Năm bảng chéo giữa cột trả lời gốc và cột mã; dưới mỗi bảng in ra “Mỗi lựa chọn chỉ ứng với một mã: True”. "
     "Đối chiếu tên lựa chọn với Bảng 3.1. Ô này không lưu file."),
    ("Ô 3. Bảng kiểm tra chất lượng dữ liệu", "Mục 3.5.1 (Bảng 3.3)",
     "Lưu bang/bang_3_5_1.csv. Bảng in ra: V2 996 0 0 6 33 204; Original 791 0 0 2 33 183. "
     "In thêm “Số dòng V2 có ít nhất một cờ: 216” và “Số dòng V2 nhiều hơn Original: 205”."),
    ("Ô 4. So khớp 791 dòng Original bằng mã băm MD5", "Mục 3.5.1",
     "In ra “Dong Original tim thay trong V2: 791 / 791” và “Cot present_in_original = True: 791”."),
    ("Ô 5. Bảng tần số 7 biến", "Mục 3.5.2 (Bảng 3.4) và Bảng 3.2 của mục 3.4",
     "Lưu bang/bang_3_5_2.csv (30 dòng, các cột bien, muc, so_luong, ty_le_pct). "
     "Ví dụ chất lượng giấc ngủ: 290, 166, 75, 214, 251 người."),
    ("Ô 6. Hình 3.3", "Mục 3.5.2",
     "Lưu hinh/hinh_3_5_2.png: 5 biểu đồ cột đen trắng, có số người trên mỗi cột."),
    ("Ô 7. Các số dùng trong phần nhận xét", "Mục 3.5.2 và Bảng 3.2",
     "In ra 45.8 %, 46.7 %, trung vị 3.0; mức 4–5: 83.2 %, 84.1 %, 83.1 %; căng thẳng mức 3–4: 92.6 %; "
     "số người ở mức 1: 17, 10, 8, 11; giới tính Male 691 (69.4), Female 305 (30.6); "
     "năm học 22, 74, 419, 481."),
]


def khoi_code(code):
    """Khung xám một ô, mỗi dòng code là một đoạn Courier New 8,5, giữ nguyên thụt lề."""
    t = tw.doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = tw.doc.styles["Table Grid"]
    c = t.cell(0, 0)
    c.width = Cm(16)
    tcpr = c._tc.get_or_add_tcPr()
    tcpr.append(tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "F2F2F2"}))
    dong = code.rstrip("\n").split("\n")
    for i, d in enumerate(dong):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent, pf.line_spacing = Cm(0), 1.0
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(4) if i == len(dong) - 1 else Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(d)
        r.font.name, r.font.size = "Courier New", Pt(8.5)
        rpr = r._element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), "Courier New")
        rpr.rFonts.set(qn("w:cs"), "Courier New")
    sau = tw.doc.add_paragraph()
    sau.paragraph_format.space_after = Pt(2)


tw.doc = tw.moi_doc()
tw.tieu_de("Code Google Colab – phần của B (Chương 3)", 2)
tw.doan("Trong phần của B, hai mục chạy số trên Google Colab là 3.5.1 (kiểm tra chất lượng dữ liệu) và 3.5.2 (mô tả "
        "dữ liệu). Mục 3.3 là phần viết, chỉ dùng thêm một ô để đối chiếu cách mã hóa với Bảng 3.1. Code dưới đây "
        "giống hệt notebook B_3_5_1_3_5_2.ipynb. Cách chạy: mở Colab, dán từng ô theo đúng thứ tự (hoặc mở thẳng "
        "notebook), chạy Ô 1 trước. Thư mục NCKH_giac_ngu trên Drive phải có cleaned_sleep_dataset.csv và "
        "cleaned_sleep_dataset_original.csv.")
tw.chu_thich("Bảng. Các mục dùng Google Colab trong phần của B")
tw.bang([
    ["Mục", "Ô code", "Thư viện", "Kết quả lưu ra"],
    ["Khởi tạo (dùng chung)", "Ô 1", "pandas, NumPy, matplotlib, SciPy, statsmodels", "Đọc 2 tệp CSV, tạo thư mục bang, hinh, ghi_chu"],
    ["3.3 Thang đo", "Ô 2", "pandas", "Không lưu file; in bảng chéo để kiểm Bảng 3.1"],
    ["3.5.1 Kiểm tra chất lượng dữ liệu", "Ô 3, Ô 4", "pandas, hashlib", "bang/bang_3_5_1.csv (Bảng 3.3)"],
    ["3.5.2 Mô tả dữ liệu", "Ô 5, Ô 6, Ô 7", "pandas, matplotlib",
     "bang/bang_3_5_2.csv (Bảng 3.4, Bảng 3.2), hinh/hinh_3_5_2.png (Hình 3.3)"],
], [4.0, 2.2, 4.4, 5.4], can_giua_tu=1, can_trai={2, 3})
tw.ghi_nguon("Ghi chú: Ô 1 nạp cả SciPy và statsmodels vì là ô khởi tạo chung của cả nhóm; phần của B không dùng "
             "hai thư viện này.")

for (tieu, muc, mong), code in zip(MO_TA, O_CODE):
    tw.tieu_de(tieu, 3)
    tw.doan(("Dùng cho: ", "dam"), muc, thut=False, keep=True)
    khoi_code(code)
    tw.doan(("Kết quả mong đợi: ", "dam"), mong, thut=False)

ra = GOC / "ch3_B_code_colab.docx"
tw.doc.save(ra)
print("đã ghi", ra)

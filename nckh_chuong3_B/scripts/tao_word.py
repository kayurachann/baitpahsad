"""Tạo phần Word của người B.

  ch3_3_3_3_5_1_3_5_2.docx: mục 3.3, 3.5.1, 3.5.2 (bản hoàn chỉnh để gửi A ráp chương), trang cuối là
  số liệu Bảng 3.2 chuyển cho C (mục 3.4).

Mọi số của 3.5.2 đọc từ ../bang/bang_3_5_2.csv. Số của tệp chính ở 3.5.1 lấy từ lần chạy thật trong
notebook của nhóm; các số của tệp Original lấy từ kết quả chạy thử trong file kế hoạch.
Đặt TO_VANG = True để tô vàng các số của tệp Original (bản nháp chờ chạy lại notebook).
Cách chạy: python3 tao_word.py
"""
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GOC = Path(__file__).resolve().parent.parent
FONT = "Times New Roman"
TO_VANG = False
NGUON = "Nguồn: Nhóm nghiên cứu tính toán từ bộ dữ liệu của Abdullah & Rahman (phiên bản 3, Mendeley Data)."

# ---------------------------------------------------------------- số liệu
b2 = pd.read_csv(GOC / "bang" / "bang_3_5_2.csv")
N = int(b2[b2["bien"] == "Giới tính"]["so_luong"].sum())


def dem(bien):
    d = b2[b2["bien"] == bien]
    return dict(zip(d["muc"].astype(str), d["so_luong"].astype(int)))


def so(x, le=1):
    """Định dạng số kiểu Việt Nam: dấu phẩy thập phân."""
    return f"{x:.{le}f}".replace(".", ",")


def pct(bien, muc):
    d = dem(bien)
    return so(round(100 * sum(d[str(m)] for m in muc) / N, 1))


def nguoi(bien, muc):
    d = dem(bien)
    return sum(d[str(m)] for m in muc)


SQ, TB, CF, HD, CT = ("Chất lượng giấc ngủ", "Thiết bị trước giờ ngủ", "Caffeine", "Hoạt động thể chất",
                      "Căng thẳng học tập")
YEU_TO = [TB, CF, HD, CT]
muc1 = [dem(b)["1"] for b in YEU_TO]
pct1 = [round(100 * dem(b)["1"] / N, 1) for b in YEU_TO]

# ---------------------------------------------------------------- định dạng
def dat_font(style_or_run, size=13, bold=None, italic=None):
    f = style_or_run.font
    f.name, f.size = FONT, Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    el = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), FONT)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(k), None)


def moi_doc():
    """Văn bản mới: A4, lề trái 3 cm, các lề khác 2 cm, Times New Roman 13, giãn dòng 1,5."""
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)
    normal = d.styles["Normal"]
    dat_font(normal)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_before, pf.space_after = 1.5, Pt(0), Pt(6)
    for ten_style, size in (("Heading 2", 13), ("Heading 3", 13)):
        st = d.styles[ten_style]
        dat_font(st, size, bold=True, italic=(ten_style == "Heading 3"))
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.keep_with_next = True
    return d


doc = moi_doc()


def doan(*phan, thut=True, canh=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, italic=False, keep=False):
    """Một đoạn văn. Mỗi phần là chuỗi, hoặc (chuỗi, 'vang') để tô vàng, hoặc (chuỗi, 'dam'/'nghieng')."""
    p = doc.add_paragraph()
    p.alignment = canh
    if thut:
        p.paragraph_format.first_line_indent = Cm(1)
    if keep:
        p.paragraph_format.keep_with_next = True
    for x in phan:
        t, kieu = (x, "") if isinstance(x, str) else x
        r = p.add_run(t)
        if size:
            r.font.size = Pt(size)
        r.italic = italic or kieu == "nghieng"
        if kieu == "dam":
            r.bold = True
        if kieu == "vang" and TO_VANG:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def chu_thich(text, tren=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if tren else 2)
    p.paragraph_format.space_after = Pt(4 if tren else 10)
    p.paragraph_format.keep_with_next = tren
    nhan, _, noi_dung = text.partition(". ")
    r = p.add_run(nhan + ". ")
    r.bold = True
    r2 = p.add_run(noi_dung)
    r2.bold = True
    return p


def ghi_nguon(text, keep=False):
    """Dòng ghi chú/nguồn dưới bảng, hình. keep=True giữ dòng này cùng trang với dòng ngay sau."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(2), Pt(10)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.italic, r.font.size = True, Pt(11)
    return p


def bang(dong, do_rong, can_giua_tu=1, o_vang=(), can_trai=(), co_chu=None, giu_lien=True):
    """Bảng kẻ ô đơn giản. dong[0] là tiêu đề. o_vang: các ô (i, j) cần tô vàng.
    Cột j < can_giua_tu hoặc j trong can_trai canh trái; co_chu: {cột: cỡ chữ} cho cột cần chữ nhỏ hơn.
    giu_lien=False cho phép bảng dài sang trang (dòng tiêu đề được lặp lại)."""
    co_chu = co_chu or {}
    t = doc.add_table(rows=len(dong), cols=len(dong[0]))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    tblpr.append(tblpr.makeelement(qn("w:tblLayout"), {qn("w:type"): "fixed"}))
    for j, w in enumerate(do_rong):
        t.columns[j].width = Cm(w)
    for i, hang in enumerate(dong):
        tr = t.rows[i]._tr
        trpr = tr.get_or_add_trPr()
        trpr.append(trpr.makeelement(qn("w:cantSplit"), {}))
        if i == 0:
            trpr.append(trpr.makeelement(qn("w:tblHeader"), {}))
        for j, gt in enumerate(hang):
            c = t.cell(i, j)
            c.width = Cm(do_rong[j])
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.0, Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.first_line_indent = Cm(0)
            giua = i == 0 or (j >= can_giua_tu and j not in can_trai)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if giua else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(gt))
            r.font.size = Pt(co_chu.get(j, 12) if i else 12)
            r.bold = i == 0
            if (i, j) in o_vang and TO_VANG:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            # các dòng giữ liền nhau; dòng cuối giữ cùng trang với ghi chú/nguồn ngay dưới bảng
            if giu_lien or i == 0 or i == len(dong) - 1:
                p.paragraph_format.keep_with_next = True
    return t


def tieu_de(text, muc):
    doc.add_heading(text, level=muc)


def tao_ban_chinh():
    """Dựng và lưu ch3_3_3_3_5_1_3_5_2.docx."""
    # ---------------------------------------------------------------- ghi chú đầu file (chỉ ở bản nháp)
    if TO_VANG:
        doan(("Ghi chú cho nhóm (xóa đoạn này khi ráp chương): ", "dam"),
         "các chỗ ", ("tô vàng", "vang"), " là số của tệp Original lấy từ kết quả chạy thử trong file kế hoạch, vì "
         "notebook hiện tại đọc nhầm tệp. Chạy lại notebook B_3_5_1_3_5_2.ipynb; nếu số in ra giống thì bỏ tô vàng, "
         "nếu khác thì sửa theo số in ra. Các số còn lại đã lấy từ lần chạy thật trên tệp chính.",
         thut=False, size=11, italic=True)

    # ================================================================= 3.3
    tieu_de("3.3. Xây dựng thang đo và bảng hỏi", 2)
    doan("Nghiên cứu dùng dữ liệu thứ cấp nên nhóm không tự thiết kế bảng hỏi. Thang đo mô tả trong mục này lấy từ "
         "bảng hỏi gốc của tác giả bộ dữ liệu (Abdullah & Rahman, phiên bản 3, Mendeley Data); bảng hỏi gồm 15 câu, "
         "ký hiệu từ Q1 đến Q15. Nhóm giữ nguyên nội dung câu hỏi và cách trả lời, chỉ chọn ra những câu phù hợp với "
         "câu hỏi nghiên cứu và mô tả lại cách mã hóa.")
    doan("Trong 15 câu, nhóm dùng 7 câu. Cả bảy câu đều là câu hỏi đóng: người trả lời chọn một phương án trong danh "
         "sách có sẵn. Câu Q6 đo chất lượng giấc ngủ và được dùng làm biến phụ thuộc. Bốn câu Q11, Q12, Q13 và Q14 đo "
         "bốn yếu tố được xem xét: sử dụng thiết bị trước giờ ngủ, tiêu thụ caffeine, hoạt động thể chất và căng thẳng "
         "học tập. Hai câu Q1 (năm học) và Q2 (giới tính) được dùng làm biến kiểm soát. Các câu này được chọn vì đo "
         "đúng chất lượng giấc ngủ và bốn yếu tố mà câu hỏi nghiên cứu đề cập, cùng hai đặc điểm của người trả lời "
         "cần kiểm soát; mỗi câu ứng với đúng một biến trong khung biến ở mục 3.2 (Hình 3.2). Bảng 3.1 trình bày tên "
         "biến, ký hiệu câu hỏi, các lựa chọn trả lời kèm mã số và tên cột tương ứng trong tệp dữ liệu.")
    chu_thich("Bảng 3.1. Thang đo các biến")
    TAN_SUAT = "Không bao giờ (1), 1–2 lần/tuần (2), 3–4 lần/tuần (3), 5–6 lần/tuần (4), "
    bang([
        ["Biến", "Câu", "Lựa chọn trả lời (mã)", "Cột trong tệp CSV"],
        ["Chất lượng giấc ngủ", "Q6", "Rất tệ (1), Tệ (2), Trung bình (3), Tốt (4), Rất tốt (5)", "sleep_quality_ordinal"],
        ["Sử dụng thiết bị trước giờ ngủ", "Q11", TAN_SUAT + "Mỗi đêm (5)", "bedtime_device_frequency_ordinal"],
        ["Tiêu thụ caffeine", "Q12", TAN_SUAT + "Mỗi ngày (5)", "caffeine_frequency_ordinal"],
        ["Hoạt động thể chất", "Q13", TAN_SUAT + "Mỗi ngày (5)", "physical_activity_frequency_ordinal"],
        ["Căng thẳng học tập", "Q14", "Không căng thẳng (1), Thấp (2), Cao (3), Rất cao (4)", "academic_stress_ordinal"],
        ["Giới tính", "Q2", "Nam, Nữ (trong tệp: Male, Female)", "gender"],
        ["Năm học", "Q1", "Năm 1, Năm 2, Năm 3, Sau đại học (trong tệp: First year, Second year, Third year, "
                          "Graduate student)", "study_year"],
    ], [3.1, 1.4, 5.9, 5.6], can_giua_tu=1, can_trai={2, 3}, co_chu={3: 10}, giu_lien=False)
    ghi_nguon("Nguồn: Nhóm nghiên cứu tổng hợp từ bảng hỏi và tệp dữ liệu của Abdullah & Rahman (phiên bản 3, "
              "Mendeley Data).")
    doan("Trong tệp dữ liệu, các lựa chọn trả lời đã được mã hóa thành số theo thứ tự từ thấp đến cao: từ 1 đến 5 với "
         "Q6, Q11, Q12 và Q13; từ 1 đến 4 với Q14, vì câu này chỉ có bốn lựa chọn. Với bốn yếu tố, số càng cao nghĩa là "
         "hành vi diễn ra càng thường xuyên (Q11, Q12, Q13) hoặc mức căng thẳng càng cao (Q14); riêng với Q6, số càng "
         "cao nghĩa là ngủ càng tốt. Mã số nằm trong các cột có đuôi _ordinal và được nhóm dùng trực tiếp. Trong tệp "
         "chính, mỗi cột mã chỉ nhận các giá trị từ 1 đến 5 (từ 1 đến 4 với căng thẳng học tập), đúng bằng số lựa chọn "
         "của câu tương ứng (xem Bảng 3.4).")
    doan("Các mã số này chỉ thể hiện thứ tự, không thể hiện khoảng cách. Chẳng hạn, khoảng cách giữa “Không bao giờ” "
         "và “1–2 lần/tuần” không nhất thiết bằng khoảng cách giữa “5–6 lần/tuần” và “Mỗi đêm”; ở Q14 cũng không có "
         "lựa chọn trung gian giữa “Thấp” và “Cao”. Vì vậy, nhóm xem các biến này là thang thứ bậc, không xem là thang "
         "khoảng cách đều. Các phép phân tích chính dùng phương pháp phù hợp với dữ liệu thứ bậc: hệ số tương quan hạng "
         "Spearman (mục 3.5.3), kiểm định Kruskal–Wallis (mục 3.5.4) và hồi quy logit thứ bậc, trong đó chất lượng "
         "giấc ngủ được xử lý như biến thứ bậc (mục 3.5.5).")
    doan("Ở một số bước khác, mã số được dùng như một điểm số: mục 3.5.4 báo cáo điểm giấc ngủ trung bình của từng "
         "nhóm; hồi quy tuyến tính (OLS) dùng mã của chất lượng giấc ngủ để ước lượng phần khác biệt mà bốn yếu tố cùng "
         "giải thích và để so sánh độ lớn liên hệ giữa các yếu tố (mục 3.5.5 và 3.5.6), đồng thời là một trong hai cách "
         "đo khi kiểm tra độ vững (mục 3.5.7). Ngoài ra, trong các mô hình hồi quy, mã của bốn yếu tố được chuẩn hóa rồi "
         "đưa vào như một biến số. Các cách dùng này đều là giả định đơn giản hóa.")
    doan("Thang đo có ba hạn chế. Thứ nhất, mỗi khái niệm chỉ được đo bằng một câu tự báo cáo. Hệ số Cronbach’s alpha "
         "cần ít nhất hai câu cùng đo một khái niệm, nên nhóm không tính được hệ số này và không đánh giá được độ tin "
         "cậy nhất quán nội tại của thang đo. Thứ hai, Q11, Q12 và Q13 hỏi về số lần trong tuần chứ không hỏi về thời "
         "lượng hay liều lượng: Q11 không cho biết mỗi lần sử dụng thiết bị trước giờ ngủ kéo dài bao lâu, Q12 không "
         "cho biết lượng caffeine tiêu thụ, Q13 không cho biết thời gian và cường độ của mỗi lần hoạt động thể chất. "
         "Hai người cùng chọn “5–6 lần/tuần” vẫn có thể khác nhau nhiều về mức độ thực tế. Thứ ba, mọi câu trả lời đều "
         "do sinh viên tự báo cáo, nên có thể sai lệch do nhớ không chính xác hoặc do xu hướng chọn câu trả lời được "
         "xem là tích cực. Các hạn chế này được nêu lại ở phần Hạn chế.")
    doan("Tám câu còn lại không được dùng làm biến. Ba câu Q3 (khó vào giấc), Q4 (thời lượng ngủ) và Q5 (thức giấc "
         "ban đêm) mô tả chính các biểu hiện của giấc ngủ, tức là gần với nội dung của biến phụ thuộc; nếu đưa các câu "
         "này vào làm yếu tố, mối liên hệ tìm được (nếu có) có thể chỉ phản ánh việc chúng và Q6 cùng đo giấc ngủ, nên "
         "kết quả không trả lời được câu hỏi nghiên cứu. Năm câu Q7 (khó tập trung), Q8 (mệt mỏi ban ngày), Q9 (vắng "
         "học), Q10 (việc hoàn thành bài tập) và Q15 (kết quả học tập) mô tả những gì có thể là hệ quả của giấc ngủ "
         "kém. Các câu này đứng sau giấc ngủ trong chuỗi giải thích, nên không phù hợp để làm yếu tố đứng trước chất "
         "lượng giấc ngủ.")
    doan("Tóm lại, năm khái niệm (chất lượng giấc ngủ và bốn yếu tố) đều được đo bằng một câu tự báo cáo theo thang "
         "thứ bậc; việc dùng một câu cho mỗi khái niệm là một hạn chế về đo lường, được nêu ở phần Hạn chế.")

    # ================================================================= 3.5.1
    tieu_de("3.5.1. Kiểm tra chất lượng dữ liệu", 3)
    doan("Trước khi phân tích, nhóm kiểm tra dữ liệu ở bốn điểm: số ô thiếu, số dòng trùng hoàn toàn (giống nhau ở cả "
         "28 cột của tệp, kể cả mốc thời gian nộp), số dòng bị gắn cờ nghi vấn và quan hệ giữa hai tệp dữ liệu. Ba loại "
         "cờ là các cột có sẵn trong dữ liệu, đánh dấu những phiếu bị lặp mốc thời gian nộp (flag_repeated_timestamp), "
         "có đáp án giống hệt (flag_identical_answers) hoặc gần trùng với phiếu khác (flag_near_duplicate). Để kiểm tra "
         "quan hệ giữa hai tệp, mỗi dòng được chuyển thành một mã băm MD5 tạo từ 15 cột trả lời gốc; sau đó, nhóm kiểm "
         "tra mã của từng dòng thuộc tệp Original có xuất hiện trong tập mã của tệp chính hay không. Các bước này được "
         "thực hiện bằng Python (thư viện pandas và hashlib) trên Google Colab. Đây là bước mô tả, không có kiểm định "
         "thống kê. Kết quả được trình bày ở Bảng 3.3.")
    chu_thich("Bảng 3.3. Kết quả kiểm tra chất lượng dữ liệu")
    bang([
        ["Tệp", "Số dòng", "Ô thiếu", "Dòng trùng hoàn toàn", "Lặp mốc thời gian", "Đáp án giống hệt", "Gần trùng"],
        ["Tệp chính", 996, 0, 0, 6, 33, 204],
        ["Tệp Original", 791, 0, 0, 2, 33, 183],
    ], [3.2, 1.8, 1.7, 2.6, 2.4, 2.3, 2.0], can_giua_tu=1, o_vang={(2, j) for j in range(1, 7)})
    ghi_nguon(NGUON)
    doan(("Cả hai tệp", "vang"), " đều không có ô thiếu và không có dòng trùng hoàn toàn. Trong tệp chính, có 6 dòng "
         "lặp mốc thời gian, 33 dòng có đáp án giống hệt và 204 dòng gần trùng; tính chung, ",
         ("216 dòng", "vang"), " bị gắn ít nhất một cờ. Cờ chỉ đánh dấu dòng đáng nghi, không nhất thiết là phiếu sai: "
         "hai sinh viên vẫn có thể nộp phiếu cùng lúc, hoặc tình cờ chọn các phương án giống nhau, nhất là khi câu hỏi "
         "có ít lựa chọn (bảy câu dùng trong bài có từ hai đến năm lựa chọn, xem Bảng 3.1). Vì vậy, nhóm giữ nguyên các "
         "dòng bị gắn cờ trong phân tích chính. Để xem kết luận có thay đổi hay không, mục 3.5.7 chạy lại mô hình theo "
         "ba cách: bỏ các dòng có đáp án giống hệt; bỏ các dòng có đáp án giống hệt hoặc gần trùng; bỏ các dòng lặp "
         "mốc thời gian.")
    doan("Về quan hệ giữa hai tệp, ", ("toàn bộ 791 dòng của tệp Original đều được tìm thấy trong tệp chính", "vang"),
         "; kết quả này khớp với cột present_in_original (791 dòng được đánh dấu). Tệp chính có thêm ",
         ("205", "vang"), " dòng so với tệp Original. Hai tệp không độc lập với nhau nên nhóm không gộp chúng lại; nếu "
         "gộp, 791 quan sát sẽ bị đếm hai lần. Tệp Original chỉ được dùng để kiểm tra độ vững ở mục 3.5.7.")
    doan("Kết luận: dữ liệu không có ô thiếu và không có dòng trùng hoàn toàn. Tệp chính có ", ("216", "vang"),
         " dòng bị gắn cờ nghi vấn; các dòng này được giữ lại và kiểm tra ở mục 3.5.7. Tất cả ", ("791", "vang"),
         " dòng của tệp Original đều nằm trong tệp chính nên hai tệp không được gộp chung.")

    # ================================================================= 3.5.2
    tieu_de("3.5.2. Mô tả dữ liệu", 3)
    doan(f"Mục này mô tả số người ở từng mức của năm biến thứ bậc trong tệp chính (n = {N}). Mục đích là xác định dữ "
         "liệu tập trung ở những mức nào, mức nào có quá ít người và chất lượng giấc ngủ phân bố như thế nào. Kết quả "
         "được trình bày ở Bảng 3.4 và Hình 3.3. Số liệu về giới tính và năm học được đặt ở mục 3.4 (Bảng 3.2).")
    chu_thich(f"Bảng 3.4. Số người (tỉ lệ %) ở từng mức của năm biến thứ bậc (n = {N})")
    ten_dong = {SQ: "Chất lượng giấc ngủ", TB: "Sử dụng thiết bị trước giờ ngủ", CF: "Tiêu thụ caffeine",
                HD: "Hoạt động thể chất", CT: "Căng thẳng học tập"}
    dong = [["Biến", "Mức 1", "Mức 2", "Mức 3", "Mức 4", "Mức 5"]]
    for b, nhan in ten_dong.items():
        d = b2[b2["bien"] == b]
        o = [f"{n} ({so(t)})" for n, t in zip(d["so_luong"], d["ty_le_pct"])]
        dong.append([nhan] + o + ["–"] * (5 - len(o)))
    bang(dong, [4.6, 2.3, 2.3, 2.3, 2.3, 2.3], can_giua_tu=1)
    hd = dem(HD)
    ghi_nguon("Ghi chú: Căng thẳng học tập chỉ có bốn mức. Với chất lượng giấc ngủ, mức 1 là rất tệ và mức 5 là rất "
              "tốt; với bốn yếu tố, mức càng cao nghĩa là càng thường xuyên hoặc càng căng thẳng. Tỉ lệ được làm tròn "
              "đến một chữ số thập phân nên tổng một dòng có thể là 99,9%; tỉ lệ gộp hai mức ở phần nhận xét được tính "
              "trực tiếp từ số người nên có thể chênh 0,1 điểm phần trăm so với cộng hai tỉ lệ trong bảng (chẳng hạn "
              f"hoạt động thể chất mức 4–5 là {hd['4'] + hd['5']}/{N} = {pct(HD, [4, 5])}%).", keep=True)
    ghi_nguon(NGUON)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(GOC / "hinh" / "hinh_3_5_2.png"), width=Cm(16))
    chu_thich(f"Hình 3.3. Phân bố số người theo mức của năm biến thứ bậc (n = {N})", tren=False)
    ghi_nguon(NGUON)

    sq = dem(SQ)
    doan(f"Chất lượng giấc ngủ có phân bố tập trung ở hai đầu thang đo. Có {nguoi(SQ, [1, 2])} người "
         f"({pct(SQ, [1, 2])}%) đánh giá giấc ngủ của mình tệ hoặc rất tệ (mức 1–2) và {nguoi(SQ, [4, 5])} người "
         f"({pct(SQ, [4, 5])}%) đánh giá tốt hoặc rất tốt (mức 4–5), trong khi mức “Trung bình” (mức 3) chỉ có "
         f"{sq['3']} người ({pct(SQ, [3])}%). Trung vị bằng 3, nhưng con số này không đại diện cho câu trả lời thường "
         "gặp vì rất ít người chọn mức giữa.")
    doan(f"Ngược lại, bốn yếu tố đều tập trung ở mức cao. Tỉ lệ người ở mức 4–5 lần lượt là {pct(TB, [4, 5])}% (sử "
         f"dụng thiết bị trước giờ ngủ), {pct(CF, [4, 5])}% (tiêu thụ caffeine) và {pct(HD, [4, 5])}% (hoạt động thể "
         f"chất); riêng căng thẳng học tập, tỉ lệ ở mức 3–4 (cao và rất cao) là {pct(CT, [3, 4])}%. Mức 1 của bốn yếu "
         f"tố chỉ có từ {min(muc1)} đến {max(muc1)} người ({so(min(pct1))}% đến {so(max(pct1))}%). Với số người ít như "
         "vậy, so sánh riêng từng mức sẽ cho kết quả không ổn định ở các nhóm nhỏ; vì thế, ở mục 3.5.4 các mức thấp "
         "được gộp lại (mức 1–3 với ba yếu tố đo bằng tần suất, mức 1–2 với căng thẳng học tập).")
    doan(f"Một số tỉ lệ cao cần được diễn giải thận trọng. Chẳng hạn, {pct(HD, [4, 5])}% sinh viên cho biết mình "
         "hoạt động thể chất từ 5 lần mỗi tuần trở lên. Vì đây là dữ liệu tự báo cáo, con số này có thể phản ánh cách "
         "người trả lời hiểu câu hỏi hoặc xu hướng chọn câu trả lời được xem là tích cực; điểm này được nêu ở phần "
         "Hạn chế.")
    doan(f"Kết luận: chất lượng giấc ngủ có phân bố tập trung ở hai đầu thang đo ({pct(SQ, [1, 2])}% tệ hoặc rất tệ, "
         f"{pct(SQ, [4, 5])}% tốt hoặc rất tốt). Bốn yếu tố tập trung ở mức cao, ít người ở mức thấp, nên phần so sánh "
         "nhóm phải gộp các mức thấp.")

    # ================================================================= Bảng 3.2 (đặt ở mục 3.4, chuyển cho C)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    doan(("Số liệu cho Bảng 3.2 (đặt ở mục 3.4)", "dam"), thut=False)
    doan("Số liệu dưới đây lấy từ cùng bảng tần số của mục 3.5.2 và được đặt ở mục 3.4.", thut=False)
    chu_thich(f"Bảng 3.2. Đặc điểm mẫu (n = {N})")
    gt, nh = dem("Giới tính"), dem("Năm học")
    b32 = [["Đặc điểm", "Nhóm", "Số lượng", "Tỉ lệ (%)"]]
    for nhom, k in (("Nam", "Male"), ("Nữ", "Female")):
        b32.append(["Giới tính", nhom, gt[k], so(round(100 * gt[k] / N, 1))])
    for nhom, k in (("Năm 1", "First year"), ("Năm 2", "Second year"), ("Năm 3", "Third year"),
                    ("Sau đại học", "Graduate student")):
        b32.append(["Năm học", nhom, nh[k], so(round(100 * nh[k] / N, 1))])
    bang(b32, [3.5, 3.5, 3.0, 3.0], can_giua_tu=2)
    ghi_nguon(NGUON)
    doan(f"Năm 1 chỉ có {nh['First year']} người nên không so sánh riêng theo năm học.", thut=False)

    ra = GOC / "ch3_3_3_3_5_1_3_5_2.docx"
    doc.save(ra)
    print("đã ghi", ra)


if __name__ == "__main__":
    tao_ban_chinh()

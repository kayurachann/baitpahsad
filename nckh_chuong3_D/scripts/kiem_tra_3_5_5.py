"""Đối chiếu số trong ch3_3_5_5_mohinh.docx với mục 3.5.5 của file kế hoạch.

Mọi số thập phân trong file Word phải có trong đoạn số liệu 3.5.5 của file kế hoạch (Bảng 3.7, OLS, VIF,
tương quan), hoặc là số tính ra từ đó: % thay đổi odds (OR − 1 hoặc 1 − OR) và OR làm tròn 2 chữ số.
Không có "ảnh hưởng", "tác động", "gây ra".
Cách chạy: python3 kiem_tra_3_5_5.py
"""
import re
import sys
from pathlib import Path

from docx import Document

GOC = Path(__file__).resolve().parent.parent
KE_HOACH = """
Thiết bị trước giờ ngủ | 1,058 (0,300) | 1,001 (0,985) | 0,991 (0,895)
Caffeine | 1,003 (0,953) | 0,884 (0,078) | 0,879 (0,069)
Hoạt động thể chất | 1,191 (0,002) | 1,261 (0,001) | 1,262 (0,001)
Căng thẳng học tập | 1,074 (0,197) | 1,015 (0,814) | 1,005 (0,933)
Nam (so với nữ) |  |  | 0,621 (0,0002)
McFadden R² |  | 0,004 | 0,014
Các mức năm học không có ý nghĩa (p từ 0,12 đến 0,74, mốc là năm 1, chỉ 22 người). OLS: R² = 0,013 (R² hiệu chỉnh
0,009); β chuẩn hóa: thiết bị 0,006; caffeine −0,073 (p = 0,056); hoạt động thể chất 0,130 (p = 0,0008); căng thẳng
0,006. VIF: 1,58; 1,61; 1,57; 1,34. Tương quan: 0,45; 0,42; 0,41; 0,40; 0,36; 0,33.
Bốn yếu tố chỉ giải thích khoảng 1,3% (OLS); OR 1,19 đến 1,26, p ≤ 0,002; OR 0,88; OR 0,62; 0,05 (ngưỡng p).
"""
hop_le = set(re.findall(r"\d+,\d+", KE_HOACH))
for x in ("1,191", "1,262", "0,621"):          # % thay đổi odds suy ra từ OR
    v = float(x.replace(",", "."))
    hop_le.add(str(round(abs(v - 1) * 100)))
hop_le.add("1,26")

doc = Document(GOC / "ch3_3_5_5_mohinh.docx")
chu = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
    c.text for t in doc.tables for r in t.rows for c in r.cells)
loi = [f"số {x} không có trong mục 3.5.5 của file kế hoạch" for x in re.findall(r"\d+,\d+", chu) if x not in hop_le]
for x in re.findall(r"khoảng (\d+)% đến (\d+)%|thấp hơn khoảng (\d+)%", chu):
    for y in filter(None, x):
        if y not in hop_le:
            loi.append(f"% thay đổi odds {y} không suy ra được từ OR")
for tu in ("ảnh hưởng", "tác động", "gây ra"):
    if tu in chu.lower():
        loi.append(f"còn từ '{tu}'")
print(f"Đã kiểm {len(re.findall(r'\d+,\d+', chu))} số thập phân.")
if loi:
    print("LỖI:\n  " + "\n  ".join(loi))
    sys.exit(1)
print("OK: mọi số khớp mục 3.5.5 của file kế hoạch, không có từ cấm.")
